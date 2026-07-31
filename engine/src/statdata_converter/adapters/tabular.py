from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd

from ..models import Dataset, DatasetMetadata
from ..report import ConversionReport
from .base import Adapter, AdapterError, FormatCapability

try:
    import xlrd
except ImportError:  # pragma: no cover - exercised by capability checks
    xlrd = None


class CsvAdapter(Adapter):
    capability = FormatCapability(
        id="csv",
        family="CSV",
        label="CSV / delimited text",
        extensions=(".csv", ".tsv"),
        can_read=True,
        can_write=True,
        trust="verified",
        detail="UTF-8 CSV and TSV with configurable delimiter and encoding.",
    )

    def read(
        self,
        path: Path,
        options: dict[str, Any],
        report: ConversionReport,
    ) -> Dataset:
        delimiter = options.get("delimiter")
        if delimiter is None:
            delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        frame = pd.read_csv(
            path,
            sep=delimiter,
            encoding=options.get("encoding", "utf-8"),
            parse_dates=options.get("parse_dates"),
        )
        report.add(
            "info",
            "csv.types_inferred",
            "Column storage types were inferred from delimited text.",
        )
        return Dataset(frame=frame, metadata=DatasetMetadata(table_name=path.stem))

    def write(
        self,
        dataset: Dataset,
        path: Path,
        options: dict[str, Any],
        report: ConversionReport,
    ) -> None:
        delimiter = options.get("delimiter")
        if delimiter is None:
            delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        dataset.frame.to_csv(
            path,
            sep=delimiter,
            encoding=options.get("encoding", "utf-8"),
            index=False,
            na_rep=options.get("na_rep", ""),
        )
        report.add(
            "info",
            "csv.storage_types_not_embedded",
            "CSV does not embed storage types; readers must infer them.",
        )


class ExcelAdapter(Adapter):
    capability = FormatCapability(
        id="excel_xlsx",
        family="Excel",
        label="Excel workbook",
        extensions=(".xlsx",),
        can_read=True,
        can_write=True,
        trust="verified",
        detail="Modern Excel workbooks. One selected worksheet per conversion.",
    )
    metadata_features = frozenset({"table_name"})

    def read(
        self,
        path: Path,
        options: dict[str, Any],
        report: ConversionReport,
    ) -> Dataset:
        sheet = options.get("sheet_name", 0)
        frame = pd.read_excel(path, sheet_name=sheet)
        sheet_name = str(sheet)
        if isinstance(sheet, int):
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True)
            sheet_name = workbook.sheetnames[sheet]
            if len(workbook.sheetnames) > 1:
                report.add(
                    "warning",
                    "excel.unselected_sheets",
                    "Only one worksheet was read; other worksheets were not converted.",
                )
            workbook.close()
        return Dataset(
            frame=frame,
            metadata=DatasetMetadata(table_name=sheet_name),
        )

    def write(
        self,
        dataset: Dataset,
        path: Path,
        options: dict[str, Any],
        report: ConversionReport,
    ) -> None:
        sheet_name = options.get(
            "sheet_name",
            dataset.metadata.table_name or "Data",
        )
        safe_sheet_name = str(sheet_name)[:31]
        dataset.frame.to_excel(path, index=False, sheet_name=safe_sheet_name)
        if str(sheet_name) != safe_sheet_name:
            report.add(
                "warning",
                "excel.sheet_name_truncated",
                "The worksheet name was truncated to Excel's 31-character limit.",
            )


class LegacyExcelAdapter(Adapter):
    def __init__(self) -> None:
        available = xlrd is not None
        self.capability = FormatCapability(
            id="excel_xls",
            family="Excel",
            label="Legacy Excel workbook",
            extensions=(".xls",),
            can_read=available,
            can_write=False,
            trust="experimental",
            availability="available" if available else "dependency_missing",
            detail=(
                "Reads one worksheet from historical Excel .xls workbooks."
                if available
                else "Legacy Excel reading requires the xlrd dependency."
            ),
            requires=("xlrd",),
        )
    metadata_features = frozenset({"table_name"})

    def read(
        self,
        path: Path,
        options: dict[str, Any],
        report: ConversionReport,
    ) -> Dataset:
        if xlrd is None:
            raise AdapterError("Legacy Excel reading requires xlrd.")
        sheet = options.get("sheet_name", 0)
        engine_kwargs: dict[str, Any] = {}
        if options.get("encoding"):
            engine_kwargs["encoding_override"] = str(options["encoding"])
        workbook = pd.ExcelFile(path, engine="xlrd", engine_kwargs=engine_kwargs)
        try:
            frame = workbook.parse(sheet_name=sheet)
            sheet_name = (
                workbook.sheet_names[sheet] if isinstance(sheet, int) else str(sheet)
            )
            if len(workbook.sheet_names) > 1:
                report.add(
                    "warning",
                    "excel.unselected_sheets",
                    "Only one worksheet was read; other worksheets were not converted.",
                )
        finally:
            workbook.close()
        report.add(
            "info",
            "excel.legacy_reader",
            "The historical Excel workbook was decoded with xlrd.",
        )
        return Dataset(
            frame=frame,
            metadata=DatasetMetadata(table_name=sheet_name),
        )

    def write(
        self,
        dataset: Dataset,
        path: Path,
        options: dict[str, Any],
        report: ConversionReport,
    ) -> None:
        del dataset, path, options, report
        raise AdapterError(
            "Legacy .xls writing is not provided; use modern .xlsx output."
        )
